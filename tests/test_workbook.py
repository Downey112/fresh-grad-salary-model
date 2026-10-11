import os
import re
import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula

from fgsalary import config, model, policy, workbook

ROOT = Path(__file__).resolve().parent.parent
REQUIRED_NAMES = {
    "Policy_Floor", "New_Min_Wage", "Current_Min_Wage", "Compression_Ceiling", "Compression_Alpha",
    "EPF_Employer", "EPF_Employee", "SOCSO_Employer", "SOCSO_Employee", "EIS_Employer", "EIS_Employee",
    "HRD_Levy", "CGPA_Mean", "X_Matrix", "X_NoConst", "Y_Vector", "Residuals", "XtX_Inverse",
    "Smearing_Factor", "DF_Residual", "Residual_Variance",
} | {f"Beta_{position}" for position in range(len(model.TERMS))}


@pytest.fixture(scope="module")
def built(result, sample, parameters, skpg, tmp_path_factory):
    path = tmp_path_factory.mktemp("workbook") / config.WORKBOOK_NAME
    scenarios = policy.run_scenarios(sample["salary"], parameters)
    workbook.build_workbook(result, scenarios, policy.spillover_profile(parameters), parameters, skpg, path)
    return path, load_workbook(path)


def formulas(book):
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, ArrayFormula):
                    yield sheet.title, cell.coordinate, value.text
                elif isinstance(value, str) and value.startswith("="):
                    yield sheet.title, cell.coordinate, value


def test_tabs_are_in_the_specified_order(built):
    _, book = built
    assert book.sheetnames[:4] == ["Parameters", "Data_and_Residuals", "OLS_Native_Formulas", "Summary_Dashboard"]


def test_every_named_range_exists_and_points_at_a_real_sheet(built):
    _, book = built
    assert REQUIRED_NAMES <= set(book.defined_names)
    for name, definition in book.defined_names.items():
        sheet_name = definition.attr_text.split("!")[0].strip("'")
        assert sheet_name in book.sheetnames, name


def test_every_name_used_in_a_formula_is_defined(built):
    _, book = built
    defined = set(book.defined_names)
    functions_and_literals = {"TRUE", "FALSE"}
    for sheet, cell, text in formulas(book):
        without_strings = re.sub(r'"[^"]*"', "", text)
        without_sheets = re.sub(r"'[^']+'!", "", without_strings)
        for token in re.findall(r"[A-Za-z_][A-Za-z0-9_.]*", without_sheets):
            is_function = re.search(rf"\b{re.escape(token)}\(", without_sheets)
            is_cell = re.fullmatch(r"\$?[A-Z]{1,3}\$?\d*", token) or re.fullmatch(r"[A-Z]{1,3}", token)
            if is_function or is_cell or token in functions_and_literals:
                continue
            assert token in defined, f"{sheet}!{cell} uses undefined name {token}"


def test_formulas_avoid_functions_that_need_a_modern_excel(built):
    _, book = built
    for sheet, cell, text in formulas(book):
        assert not re.search(r"T\.DIST|T\.INV|F\.DIST|\bLET\(|\bLAMBDA\(|\bFILTER\(|_xlfn", text), f"{sheet}!{cell}"


def test_parameters_hold_the_configured_values(built, parameters):
    _, book = built
    sheet = book["Parameters"]
    values = {sheet[f"C{row}"].value: sheet[f"B{row}"].value for row in range(4, 16)}
    assert values["Policy_Floor"] == parameters["policy"]["graduate_floor"]
    assert values["New_Min_Wage"] == parameters["policy"]["new_minimum_wage"]
    assert values["Compression_Ceiling"] == parameters["policy"]["compression_ceiling"]
    assert values["Compression_Alpha"] == parameters["policy"]["compression_alpha"]
    assert values["EPF_Employer"] == parameters["statutory"]["epf_employer"]
    assert values["HRD_Levy"] == parameters["statutory"]["hrd_levy"]


def test_data_tab_holds_every_graduate_with_live_formulas(built, sample):
    _, book = built
    sheet = book["Data_and_Residuals"]
    assert sheet.max_row == len(sample) + 1
    assert [cell.value for cell in sheet[1]] == workbook.DATA_HEADERS
    last = sheet.max_row
    assert sheet[f"G{last}"].value == int(sample["salary"].iloc[-1])
    assert sheet[f"N{last}"].value == f"=LN(G{last})"
    assert sheet[f"P{last}"].value == f"=N{last}-O{last}"
    assert book.defined_names["Y_Vector"].attr_text.endswith(f"$N$2:$N${last}")


def test_ols_tab_uses_linest_and_matrix_algebra(built):
    _, book = built
    sheet = book["OLS_Native_Formulas"]
    linest = sheet[f"B{workbook.OLS_LINEST_TOP}"].value
    closed_form = sheet[f"B{workbook.OLS_BETA_TOP}"].value
    assert isinstance(linest, ArrayFormula) and "LINEST(Y_Vector,X_NoConst,TRUE,TRUE)" in linest.text
    assert linest.ref == f"B{workbook.OLS_LINEST_TOP}:G{workbook.OLS_LINEST_TOP + 4}"
    assert isinstance(closed_form, ArrayFormula)
    assert closed_form.text == (
        "=MMULT(MINVERSE(MMULT(TRANSPOSE(X_Matrix),X_Matrix)),MMULT(TRANSPOSE(X_Matrix),Y_Vector))")
    assert closed_form.ref == f"B{workbook.OLS_BETA_TOP}:B{workbook.OLS_BETA_TOP + len(model.TERMS) - 1}"


def test_workbook_recalculates_when_opened(built):
    _, book = built
    assert book.calculation.fullCalcOnLoad is True


@pytest.mark.skipif(
    os.environ.get("FGSALARY_EXCEL") != "1" or shutil.which("cscript") is None,
    reason="set FGSALARY_EXCEL=1 on a Windows machine with Excel to run the live check",
)
def test_excel_itself_reproduces_python(built, result, parameters, tmp_path):
    path, _ = built
    dump = tmp_path / "excel.txt"
    subprocess.run(
        ["cscript", "//nologo", str(ROOT / "tools" / "verify_workbook.vbs"), str(path), str(dump)],
        check=True, timeout=300,
    )
    lines = dump.read_text(encoding="utf-16" if dump.read_bytes()[:2] == b"\xff\xfe" else "mbcs").splitlines()
    split = lines.index("profile=low")
    cells = dict(line[5:].split("|", 1) for line in lines[:split] if line.startswith("cell="))
    low = dict(line[5:].split("|", 1) for line in lines[split:] if line.startswith("cell="))
    for line in lines:
        if line.startswith("sheet="):
            assert line.endswith("errors=0"), line

    coefficients = result["coefficients"]["coefficient"].to_numpy()
    for position, expected in enumerate(coefficients):
        row = workbook.OLS_BETA_TOP + position
        assert float(cells[f"OLS_Native_Formulas!B{row}"]) == pytest.approx(expected, abs=1e-9)
        assert float(cells[f"OLS_Native_Formulas!H{row}"]) == pytest.approx(expected, abs=1e-9)

    policy_parameters, statutory = parameters["policy"], parameters["statutory"]
    for values, profile in ((cells, (3.5, 6, 1, 1)), (low, (2.5, 3, 0, 0))):
        base = model.predict_salary(result, *profile)
        post = float(policy.post_policy_wage(
            base, policy_parameters["graduate_floor"], policy_parameters["current_minimum_wage"],
            policy_parameters["compression_ceiling"], policy_parameters["compression_alpha"]))
        out = workbook.CALC_OUTPUT_TOP
        assert float(values[f"Summary_Dashboard!B{out + 1}"]) == pytest.approx(base, rel=1e-9)
        assert float(values[f"Summary_Dashboard!B{out + 2}"]) == pytest.approx(post, rel=1e-9)
        assert float(values[f"Summary_Dashboard!B{out + 7}"]) == pytest.approx(
            float(policy.employer_cost(post, statutory)), rel=1e-9)
        assert float(values[f"Summary_Dashboard!B{out + 9}"]) == pytest.approx(
            float(policy.take_home(post, statutory)), rel=1e-9)
