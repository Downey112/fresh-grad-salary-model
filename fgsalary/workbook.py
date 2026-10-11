"""Build the Excel workbook.

The workbook recomputes the regression with native Excel formulas, so a
reader can change a salary or a parameter and watch everything update. Values
that Excel cannot produce on its own (HC3 standard errors, VIFs, the scenario
table) are written as numbers and labelled as coming from Python.

Compatibility choices: array formulas are written as classic Ctrl+Shift+Enter
arrays and the statistical functions are the long-standing ones (TDIST, TINV,
FDIST), so the file calculates in any Excel from 2007 on without #NAME? errors.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.table import Table, TableStyleInfo

from . import data as data_module
from .model import LABELS, TERMS

SHEET_PARAMETERS = "Parameters"
SHEET_DATA = "Data_and_Residuals"
SHEET_OLS = "OLS_Native_Formulas"
SHEET_DASHBOARD = "Summary_Dashboard"
SHEET_POLICY = "Policy_Simulation"
SHEET_NOTES = "Notes"

TITLE_FONT = Font(bold=True, size=14)
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SECTION_FONT = Font(bold=True, size=11, color="1F4E78")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
OUTPUT_FILL = PatternFill("solid", fgColor="E2EFDA")
PYTHON_FILL = PatternFill("solid", fgColor="EDEDED")
NOTE_FONT = Font(italic=True, color="595959")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

MONEY = '#,##0.00'
WHOLE_MONEY = '#,##0'
PERCENT = '0.00%'
COEFFICIENT = '0.000000'
STATISTIC = '0.0000'
TINY = '0.00E+00'

# (label, named range, parameter group, parameter key, number format, note)
PARAMETER_ROWS = [
    ("Graduate wage floor (RM/month)", "Policy_Floor", "policy", "graduate_floor", WHOLE_MONEY,
     "Budget 2027: floor for graduates and semi-skilled workers. Start date not announced."),
    ("New national minimum wage (RM/month)", "New_Min_Wage", "policy", "new_minimum_wage", WHOLE_MONEY,
     "Budget 2027: effective June 2027."),
    ("Minimum wage before the policy (RM/month)", "Current_Min_Wage", "policy", "current_minimum_wage",
     WHOLE_MONEY, "In force since February 2025. Anchors the wage compression rule."),
    ("Wage compression ceiling (RM/month)", "Compression_Ceiling", "policy", "compression_ceiling",
     WHOLE_MONEY, "Assumption: pay at or above this is unaffected by the floor."),
    ("Wage compression alpha", "Compression_Alpha", "policy", "compression_alpha", '0.00',
     "Assumption: share of squeezed pay gaps that employers restore (0 = none, 1 = all)."),
    ("EPF, employer", "EPF_Employer", "statutory", "epf_employer", PERCENT,
     "Rate for monthly wages of RM5,000 and below."),
    ("EPF, employee", "EPF_Employee", "statutory", "epf_employee", PERCENT, ""),
    ("SOCSO, employer", "SOCSO_Employer", "statutory", "socso_employer", PERCENT,
     "Flat-rate approximation of the contribution table."),
    ("SOCSO, employee", "SOCSO_Employee", "statutory", "socso_employee", PERCENT,
     "Flat-rate approximation of the contribution table."),
    ("EIS, employer", "EIS_Employer", "statutory", "eis_employer", PERCENT, ""),
    ("EIS, employee", "EIS_Employee", "statutory", "eis_employee", PERCENT, ""),
    ("HRD Corp levy, employer", "HRD_Levy", "statutory", "hrd_levy", PERCENT,
     "Applies to employers with 10 or more Malaysian employees."),
]

DATA_HEADERS = [
    "graduate_id", "field", "cgpa", "intern_months", "stem", "klang_valley", "salary",
    "const", "cgpa_c", "x_intern_months", "x_stem", "x_klang_valley", "cgpa_c_x_stem",
    "ln_salary", "fitted_ln_salary", "residual", "fitted_salary",
]

OLS_LINEST_TOP = 6          # first row of the 5-row LINEST block (columns B:G)
OLS_BETA_TOP = 15           # first row of the six closed-form coefficients
OLS_INVERSE_TOP = 24        # first row of the 6 x 6 inverse (columns B:G)
OLS_STATS_TOP = 33          # first row of the fit statistics
DASHBOARD_TABLE_TOP = 5     # first coefficient row on the dashboard
CALC_INPUT_TOP = 24         # first calculator input row
CALC_OUTPUT_TOP = 31        # first calculator output row


def name(workbook, range_name, sheet, reference):
    workbook.defined_names[range_name] = DefinedName(range_name, attr_text=f"'{sheet}'!{reference}")


def header_row(sheet, row, headers, first_column=1):
    for offset, text in enumerate(headers):
        cell = sheet.cell(row=row, column=first_column + offset, value=text)
        cell.font, cell.fill, cell.border = HEADER_FONT, HEADER_FILL, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def put(sheet, reference, value, number_format=None, fill=None, bold=False):
    cell = sheet[reference]
    cell.value = value
    cell.border = BORDER
    if number_format:
        cell.number_format = number_format
    if fill:
        cell.fill = fill
    if bold:
        cell.font = Font(bold=True)
    return cell


def widths(sheet, sizes):
    for column, width in sizes.items():
        sheet.column_dimensions[column].width = width


def build_parameters(workbook, parameters, cgpa_mean):
    sheet = workbook.active
    sheet.title = SHEET_PARAMETERS
    sheet["A1"] = "Parameters"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = "Yellow cells are inputs. Every formula in the workbook reads them through the named range in column C."
    sheet["A2"].font = NOTE_FONT
    header_row(sheet, 3, ["Parameter", "Value", "Named range", "Note"])

    row = 4
    for label, range_name, group, key, number_format, note in PARAMETER_ROWS:
        put(sheet, f"A{row}", label)
        put(sheet, f"B{row}", parameters[group][key], number_format, INPUT_FILL)
        put(sheet, f"C{row}", range_name)
        put(sheet, f"D{row}", note)
        name(workbook, range_name, SHEET_PARAMETERS, f"$B${row}")
        row += 1

    put(sheet, f"A{row}", "Mean CGPA of the sample (centring constant)")
    put(sheet, f"B{row}", f"=AVERAGE('{SHEET_DATA}'!$C$2:$C${parameters['sample_size'] + 1})", '0.0000', OUTPUT_FILL)
    put(sheet, f"C{row}", "CGPA_Mean")
    put(sheet, f"D{row}", f"Formula. statsmodels used {cgpa_mean:.6f}.")
    name(workbook, "CGPA_Mean", SHEET_PARAMETERS, f"$B${row}")
    widths(sheet, {"A": 44, "B": 14, "C": 22, "D": 82})
    sheet.freeze_panes = "A4"


def build_data(workbook, frame):
    sheet = workbook.create_sheet(SHEET_DATA)
    header_row(sheet, 1, DATA_HEADERS)
    last = len(frame) + 1

    for offset, record in enumerate(frame.itertuples(index=False)):
        row = offset + 2
        sheet.cell(row=row, column=1, value=int(record.graduate_id))
        sheet.cell(row=row, column=2, value=record.field)
        sheet.cell(row=row, column=3, value=float(record.cgpa)).number_format = '0.00'
        sheet.cell(row=row, column=4, value=int(record.intern_months))
        sheet.cell(row=row, column=5, value=int(record.stem))
        sheet.cell(row=row, column=6, value=int(record.klang_valley))
        sheet.cell(row=row, column=7, value=int(record.salary)).number_format = WHOLE_MONEY
        sheet.cell(row=row, column=8, value=1)
        sheet.cell(row=row, column=9, value=f"=C{row}-CGPA_Mean").number_format = '0.0000'
        sheet.cell(row=row, column=10, value=f"=D{row}")
        sheet.cell(row=row, column=11, value=f"=E{row}")
        sheet.cell(row=row, column=12, value=f"=F{row}")
        sheet.cell(row=row, column=13, value=f"=I{row}*K{row}").number_format = '0.0000'
        sheet.cell(row=row, column=14, value=f"=LN(G{row})").number_format = COEFFICIENT
        sheet.cell(
            row=row, column=15,
            value=f"=H{row}*Beta_0+I{row}*Beta_1+J{row}*Beta_2+K{row}*Beta_3+L{row}*Beta_4+M{row}*Beta_5",
        ).number_format = COEFFICIENT
        sheet.cell(row=row, column=16, value=f"=N{row}-O{row}").number_format = COEFFICIENT
        sheet.cell(row=row, column=17, value=f"=EXP(O{row})*Smearing_Factor").number_format = MONEY

    table = Table(displayName="GraduateData", ref=f"A1:{get_column_letter(len(DATA_HEADERS))}{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    sheet.freeze_panes = "C2"
    widths(sheet, {"A": 12, "B": 42, **{get_column_letter(i): 15 for i in range(3, 18)}})

    name(workbook, "X_Matrix", SHEET_DATA, f"$H$2:$M${last}")
    name(workbook, "X_NoConst", SHEET_DATA, f"$I$2:$M${last}")
    name(workbook, "Y_Vector", SHEET_DATA, f"$N$2:$N${last}")
    name(workbook, "Residuals", SHEET_DATA, f"$P$2:$P${last}")


def build_ols(workbook, result):
    sheet = workbook.create_sheet(SHEET_OLS)
    coefficients = result["coefficients"].set_index("term")
    statistics = result["fit_statistics"]

    sheet["A1"] = "OLS with native Excel formulas"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = ("Dependent variable: ln(salary). Two independent Excel routes to the same coefficients, "
                   "each checked against statsmodels.")
    sheet["A2"].font = NOTE_FONT

    # 1. LINEST. It returns the regressors in reverse order, intercept last.
    sheet["A4"] = "1. LINEST(Y_Vector, X_NoConst, TRUE, TRUE)"
    sheet["A4"].font = SECTION_FONT
    header_row(sheet, 5, ["LINEST output"] + [LABELS[term] for term in reversed(TERMS)])
    linest_labels = [
        "Coefficients", "Standard errors", "R-squared | residual std error",
        "F statistic | residual df", "Regression SS | residual SS",
    ]
    for offset, label in enumerate(linest_labels):
        put(sheet, f"A{OLS_LINEST_TOP + offset}", label)
    block = f"B{OLS_LINEST_TOP}:G{OLS_LINEST_TOP + 4}"
    sheet[f"B{OLS_LINEST_TOP}"] = ArrayFormula(block, "=IFERROR(LINEST(Y_Vector,X_NoConst,TRUE,TRUE),\"\")")
    for row in sheet[block]:
        for cell in row:
            cell.number_format, cell.border, cell.fill = COEFFICIENT, BORDER, OUTPUT_FILL
    sheet[f"A{OLS_LINEST_TOP + 5}"] = ("Rows 3 to 5 only use the first two columns; "
                                       "LINEST returns #N/A for the rest, shown here as blanks.")
    sheet[f"A{OLS_LINEST_TOP + 5}"].font = NOTE_FONT

    # 2. Closed form: beta = (X'X)^-1 X'y.
    sheet[f"A{OLS_BETA_TOP - 2}"] = "2. Closed form: MMULT(MINVERSE(MMULT(TRANSPOSE(X),X)),MMULT(TRANSPOSE(X),Y))"
    sheet[f"A{OLS_BETA_TOP - 2}"].font = SECTION_FONT
    header_row(sheet, OLS_BETA_TOP - 1, [
        "Term", "Coefficient (MMULT)", "Std error", "t value", "p value", "95% CI low", "95% CI high",
        "LINEST coefficient", "MMULT minus LINEST", "statsmodels coefficient", "MMULT minus statsmodels",
    ])
    last_beta = OLS_BETA_TOP + len(TERMS) - 1
    sheet[f"B{OLS_BETA_TOP}"] = ArrayFormula(
        f"B{OLS_BETA_TOP}:B{last_beta}",
        "=MMULT(MINVERSE(MMULT(TRANSPOSE(X_Matrix),X_Matrix)),MMULT(TRANSPOSE(X_Matrix),Y_Vector))",
    )
    inverse = f"$B${OLS_INVERSE_TOP}:$G${OLS_INVERSE_TOP + len(TERMS) - 1}"
    for position, term in enumerate(TERMS):
        row = OLS_BETA_TOP + position
        diagonal = position + 1
        put(sheet, f"A{row}", LABELS[term])
        sheet[f"B{row}"].number_format, sheet[f"B{row}"].border, sheet[f"B{row}"].fill = (
            COEFFICIENT, BORDER, OUTPUT_FILL)
        put(sheet, f"C{row}", f"=SQRT(Residual_Variance*INDEX({inverse},{diagonal},{diagonal}))", COEFFICIENT)
        put(sheet, f"D{row}", f"=B{row}/C{row}", STATISTIC)
        put(sheet, f"E{row}", f"=TDIST(ABS(D{row}),DF_Residual,2)", STATISTIC)
        put(sheet, f"F{row}", f"=B{row}-TINV(0.05,DF_Residual)*C{row}", COEFFICIENT)
        put(sheet, f"G{row}", f"=B{row}+TINV(0.05,DF_Residual)*C{row}", COEFFICIENT)
        put(sheet, f"H{row}", f"=INDEX($B${OLS_LINEST_TOP}:$G${OLS_LINEST_TOP},1,{len(TERMS) - position})", COEFFICIENT)
        put(sheet, f"I{row}", f"=B{row}-H{row}", TINY)
        put(sheet, f"J{row}", float(coefficients.loc[term, "coefficient"]), COEFFICIENT, PYTHON_FILL)
        put(sheet, f"K{row}", f"=B{row}-J{row}", TINY)
        name(workbook, f"Beta_{position}", SHEET_OLS, f"$B${row}")

    # 3. The inverse the standard errors come from.
    sheet[f"A{OLS_INVERSE_TOP - 2}"] = "3. (X'X)^-1 = MINVERSE(MMULT(TRANSPOSE(X_Matrix),X_Matrix))"
    sheet[f"A{OLS_INVERSE_TOP - 2}"].font = SECTION_FONT
    header_row(sheet, OLS_INVERSE_TOP - 1, [""] + [LABELS[term] for term in TERMS])
    block = f"B{OLS_INVERSE_TOP}:G{OLS_INVERSE_TOP + len(TERMS) - 1}"
    sheet[f"B{OLS_INVERSE_TOP}"] = ArrayFormula(block, "=MINVERSE(MMULT(TRANSPOSE(X_Matrix),X_Matrix))")
    for position, term in enumerate(TERMS):
        put(sheet, f"A{OLS_INVERSE_TOP + position}", LABELS[term])
    for row in sheet[block]:
        for cell in row:
            cell.number_format, cell.border = '0.00000000', BORDER
    name(workbook, "XtX_Inverse", SHEET_OLS, inverse)

    # 4. Fit statistics, each next to the statsmodels value.
    top = OLS_STATS_TOP
    sheet[f"A{top - 2}"] = "4. Fit statistics"
    sheet[f"A{top - 2}"].font = SECTION_FONT
    header_row(sheet, top - 1, ["Statistic", "Excel formula", "statsmodels", "Excel minus statsmodels"])
    rows = [
        ("Observations", "=ROWS(Y_Vector)", statistics["observations"], '0', "N_Obs"),
        ("Parameters (incl. intercept)", "=COLUMNS(X_Matrix)", len(TERMS), '0', "N_Params"),
        ("Residual degrees of freedom", "=N_Obs-N_Params", statistics["df_residual"], '0', "DF_Residual"),
        ("Residual sum of squares", "=SUMSQ(Residuals)", float(result["classical"].ssr), STATISTIC, "SS_Residual"),
        ("Total sum of squares", "=DEVSQ(Y_Vector)", float(result["classical"].centered_tss), STATISTIC, "SS_Total"),
        ("Residual variance", "=SS_Residual/DF_Residual", float(result["classical"].scale), COEFFICIENT,
         "Residual_Variance"),
        ("R-squared", "=1-SS_Residual/SS_Total", statistics["r_squared"], COEFFICIENT, "R_Squared"),
        ("Adjusted R-squared", "=1-(1-R_Squared)*(N_Obs-1)/DF_Residual", statistics["adj_r_squared"],
         COEFFICIENT, "Adj_R_Squared"),
        ("F statistic", "=((SS_Total-SS_Residual)/(N_Params-1))/Residual_Variance", statistics["f_statistic"],
         STATISTIC, "F_Statistic"),
        ("F p-value", "=FDIST(F_Statistic,N_Params-1,DF_Residual)", statistics["f_p_value"], TINY, "F_P_Value"),
        ("Log-likelihood", "=-N_Obs/2*(LN(2*PI())+LN(SS_Residual/N_Obs)+1)", statistics["log_likelihood"],
         STATISTIC, "Log_Likelihood"),
        ("AIC", "=2*N_Params-2*Log_Likelihood", statistics["aic"], STATISTIC, "Model_AIC"),
        ("BIC", "=N_Params*LN(N_Obs)-2*Log_Likelihood", statistics["bic"], STATISTIC, "Model_BIC"),
        ("Smearing factor (mean of exp(residual))", "=SUMPRODUCT(EXP(Residuals))/ROWS(Residuals)",
         result["smearing"], COEFFICIENT, "Smearing_Factor"),
    ]
    for offset, (label, formula, python_value, number_format, range_name) in enumerate(rows):
        row = top + offset
        put(sheet, f"A{row}", label)
        put(sheet, f"B{row}", formula, number_format, OUTPUT_FILL)
        put(sheet, f"C{row}", python_value, number_format, PYTHON_FILL)
        put(sheet, f"D{row}", f"=B{row}-C{row}", TINY)
        name(workbook, range_name, SHEET_OLS, f"$B${row}")

    widths(sheet, {"A": 40, **{get_column_letter(i): 22 for i in range(2, 12)}})


def build_dashboard(workbook, result, parameters):
    sheet = workbook.create_sheet(SHEET_DASHBOARD)
    coefficients = result["coefficients"].set_index("term")
    top = DASHBOARD_TABLE_TOP

    sheet["A1"] = "Fresh graduate starting salary model: summary"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = ("Simulated sample calibrated to KPT's Graduate Tracer Study 2024 salary bands. "
                   "CGPA, internship and location effects reflect assumptions; see the Notes tab.")
    sheet["A2"].font = NOTE_FONT
    sheet["A3"] = "Regression summary: ln(salary)"
    sheet["A3"].font = SECTION_FONT
    header_row(sheet, top - 1, [
        "Term", "Coefficient", "Std error (classical)", "Std error (HC3)", "t (HC3)", "p (HC3)",
        "Sig.", "VIF", "Effect on salary", "Read the effect as",
    ])
    effect_units = {
        "const": "",
        "cgpa_c": "per 1.00 CGPA point, non-STEM graduate",
        "intern_months": "per extra month of internship",
        "stem": "STEM vs non-STEM at the average CGPA",
        "klang_valley": "Klang Valley vs elsewhere",
        "cgpa_c_x_stem": "extra return to 1.00 CGPA point for STEM",
    }
    for position, term in enumerate(TERMS):
        row = top + position
        put(sheet, f"A{row}", LABELS[term])
        put(sheet, f"B{row}", f"=Beta_{position}", COEFFICIENT, OUTPUT_FILL)
        put(sheet, f"C{row}", f"='{SHEET_OLS}'!C{OLS_BETA_TOP + position}", COEFFICIENT, OUTPUT_FILL)
        put(sheet, f"D{row}", float(coefficients.loc[term, "hc3_std_error"]), COEFFICIENT, PYTHON_FILL)
        put(sheet, f"E{row}", f"=IFERROR(B{row}/D{row},\"\")", STATISTIC, OUTPUT_FILL)
        put(sheet, f"F{row}", f"=IFERROR(TDIST(ABS(E{row}),DF_Residual,2),\"\")", STATISTIC, OUTPUT_FILL)
        put(sheet, f"G{row}", f"=IF(F{row}=\"\",\"\",IF(F{row}<0.01,\"***\",IF(F{row}<0.05,\"**\",IF(F{row}<0.1,\"*\",\"\"))))")
        sheet[f"G{row}"].alignment = Alignment(horizontal="center")
        if term == "const":
            put(sheet, f"H{row}", "")
            put(sheet, f"I{row}", "")
        else:
            put(sheet, f"H{row}", float(coefficients.loc[term, "vif"]), '0.00', PYTHON_FILL)
            put(sheet, f"I{row}", f"=EXP(B{row})-1", PERCENT, OUTPUT_FILL)
        put(sheet, f"J{row}", effect_units[term])
    notes_row = top + len(TERMS)
    sheet[f"A{notes_row}"] = ("*** p<0.01, ** p<0.05, * p<0.10, using HC3 heteroscedasticity-robust standard errors. "
                              "Green cells are live Excel formulas; grey cells are values computed in Python.")
    sheet[f"A{notes_row}"].font = NOTE_FONT

    fit_top = notes_row + 2
    sheet[f"A{fit_top - 1}"] = "Model fit"
    sheet[f"A{fit_top - 1}"].font = SECTION_FONT
    fit_rows = [
        ("Observations", "=N_Obs", '0'),
        ("R-squared", "=R_Squared", '0.0000'),
        ("Adjusted R-squared", "=Adj_R_Squared", '0.0000'),
        ("F statistic", "=F_Statistic", STATISTIC),
        ("AIC", "=Model_AIC", STATISTIC),
        ("BIC", "=Model_BIC", STATISTIC),
        ("Smearing factor", "=Smearing_Factor", '0.0000'),
    ]
    for offset, (label, formula, number_format) in enumerate(fit_rows):
        put(sheet, f"A{fit_top + offset}", label)
        put(sheet, f"B{fit_top + offset}", formula, number_format, OUTPUT_FILL)

    # What-if calculator.
    first = CALC_INPUT_TOP
    sheet[f"A{first - 2}"] = "What-if calculator"
    sheet[f"A{first - 2}"].font = SECTION_FONT
    header_row(sheet, first - 1, ["Candidate profile", "Input", "Allowed values"])
    inputs = [
        ("CGPA", 3.5, "2.00 to 4.00", '0.00'),
        ("Internship months", 6, "3, 6 or 9", '0'),
        ("STEM degree", 1, "1 = yes, 0 = no", '0'),
        ("Works in Klang Valley", 1, "1 = yes, 0 = no", '0'),
        ("Graduate floor applies", 1, "1 = yes, 0 = no (e.g. exempt employer)", '0'),
    ]
    for offset, (label, default, allowed, number_format) in enumerate(inputs):
        put(sheet, f"A{first + offset}", label)
        put(sheet, f"B{first + offset}", default, number_format, INPUT_FILL)
        put(sheet, f"C{first + offset}", allowed)
    cgpa, intern, stem, kv, applies = (f"$B${first + offset}" for offset in range(5))

    simulation = parameters["simulation"]
    cgpa_rule = DataValidation(type="decimal", operator="between",
                               formula1=str(simulation["cgpa_min"]), formula2=str(simulation["cgpa_max"]),
                               showErrorMessage=True, errorTitle="CGPA", error="Enter a CGPA between 2.00 and 4.00.")
    months = ",".join(simulation["intern_months_probabilities"])
    intern_rule = DataValidation(type="list", formula1=f'"{months}"', showErrorMessage=True,
                                 errorTitle="Internship months", error=f"Choose one of: {months}.")
    binary_rule = DataValidation(type="list", formula1='"0,1"', showErrorMessage=True,
                                 errorTitle="Yes or no", error="Enter 1 for yes or 0 for no.")
    for rule in (cgpa_rule, intern_rule, binary_rule):
        sheet.add_data_validation(rule)
    cgpa_rule.add(cgpa.replace("$", ""))
    intern_rule.add(intern.replace("$", ""))
    for reference in (stem, kv, applies):
        binary_rule.add(reference.replace("$", ""))

    out = CALC_OUTPUT_TOP
    header_row(sheet, out - 1, ["Forecast", "RM per month", "How it is worked out"])
    base, post, to_floor = f"$B${out + 1}", f"$B${out + 2}", f"$B${out + 3}"
    remap = ("Policy_Floor+(MIN(MAX({base},Current_Min_Wage),Compression_Ceiling)-Current_Min_Wage)"
             "*(Compression_Ceiling-Policy_Floor)/(Compression_Ceiling-Current_Min_Wage)").format(base=base)
    outputs = [
        ("Predicted ln(salary)",
         f"=Beta_0+Beta_1*({cgpa}-CGPA_Mean)+Beta_2*{intern}+Beta_3*{stem}+Beta_4*{kv}"
         f"+Beta_5*({cgpa}-CGPA_Mean)*{stem}",
         COEFFICIENT, "Regression coefficients times the profile"),
        ("Forecast baseline salary", f"=IFERROR(EXP($B${out})*Smearing_Factor,0)", MONEY,
         "exp(prediction) times the smearing factor"),
        ("Post-policy salary",
         f"=IFERROR(IF({applies}=1,IF({base}>=Compression_Ceiling,{base},"
         f"(1-Compression_Alpha)*MAX({base},Policy_Floor)+Compression_Alpha*({remap})),{base}),{base})",
         MONEY, "Floor plus partial restoration of squeezed pay gaps"),
        ("Raise needed to reach the floor", f"=IF({applies}=1,MAX(Policy_Floor-{base},0),0)", MONEY,
         "Zero if the baseline is already at or above the floor"),
        ("Wage compression adjustment", f"={post}-{base}-{to_floor}", MONEY,
         "Extra raise from restoring pay gaps (spillover)"),
        ("Total raise", f"={post}-{base}", MONEY, "Post-policy minus baseline"),
        ("Employer statutory on-cost", f"={post}*(EPF_Employer+SOCSO_Employer+EIS_Employer+HRD_Levy)", MONEY,
         "EPF + SOCSO + EIS + HRD Corp levy on post-policy salary"),
        ("Total cost of employment (TCOE)", f"={post}+$B${out + 6}", MONEY, "Salary plus employer on-cost"),
        ("Employee statutory deductions", f"={post}*(EPF_Employee+SOCSO_Employee+EIS_Employee)", MONEY,
         "EPF + SOCSO + EIS"),
        ("Net take-home pay (before income tax)", f"={post}-$B${out + 8}", MONEY,
         "Income tax (PCB) is not deducted"),
    ]
    for offset, (label, formula, number_format, explanation) in enumerate(outputs):
        put(sheet, f"A{out + offset}", label)
        put(sheet, f"B{out + offset}", formula, number_format, OUTPUT_FILL, bold=offset in (1, 2, 7, 9))
        put(sheet, f"C{out + offset}", explanation)

    widths(sheet, {"A": 40, "B": 18, "C": 46, "D": 18, "E": 12, "F": 12, "G": 8, "H": 10, "I": 16, "J": 44})


def build_policy(workbook, scenarios, profile, skpg, frame):
    sheet = workbook.create_sheet(SHEET_POLICY)
    sheet["A1"] = "Wage floor simulation (computed in Python)"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = ("Forward-looking: the floor is announced but not in force, so these are simulated outcomes "
                   "under the stated assumptions, not measured effects.")
    sheet["A2"].font = NOTE_FONT

    columns = [
        ("scenario", "Scenario", None),
        ("floor", "Floor (RM)", WHOLE_MONEY),
        ("share_below_floor", "Share below floor", PERCENT),
        ("share_receiving_raise", "Share receiving a raise", PERCENT),
        ("mean_raise_if_raised", "Mean raise if raised (RM)", MONEY),
        ("median_salary_before", "Median before (RM)", WHOLE_MONEY),
        ("median_salary_after", "Median after (RM)", WHOLE_MONEY),
        ("wage_bill_increase", "Wage bill increase", PERCENT),
        ("employer_cost_increase_per_1000_grads", "Extra employer cost per 1,000 graduates (RM/month)", WHOLE_MONEY),
        ("mid_tier_share_raised", "Mid-tier share raised", PERCENT),
        ("mid_tier_mean_raise", "Mid-tier mean raise (RM)", MONEY),
        ("mid_tier_mean_raise_pct", "Mid-tier mean raise (%)", PERCENT),
        ("p50_p10_ratio_before", "P50/P10 before", '0.00'),
        ("p50_p10_ratio_after", "P50/P10 after", '0.00'),
        ("description", "What the scenario assumes", None),
    ]
    header_row(sheet, 4, [title for _, title, _ in columns])
    for offset, record in enumerate(scenarios.to_dict("records")):
        for position, (key, _, number_format) in enumerate(columns, start=1):
            value = record[key]
            cell = sheet.cell(row=5 + offset, column=position, value=value if isinstance(value, str) else float(value))
            cell.border = BORDER
            if number_format:
                cell.number_format = number_format
    sheet.row_dimensions[4].height = 48

    top = 5 + len(scenarios) + 3
    sheet[f"A{top - 2}"] = "Raise by pre-policy salary, central case"
    sheet[f"A{top - 2}"].font = SECTION_FONT
    header_row(sheet, top - 1, ["Salary before (RM)", "Salary after (RM)", "Raise from floor (RM)",
                                "Raise from spillover (RM)", "Total raise (RM)"])
    for offset, record in enumerate(profile.itertuples(index=False)):
        for position, value in enumerate(record, start=1):
            cell = sheet.cell(row=top + offset, column=position, value=float(value))
            cell.number_format, cell.border = MONEY, BORDER

    top = top + len(profile) + 3
    sheet[f"A{top - 2}"] = "Calibration check: salary bands, published versus simulated sample"
    sheet[f"A{top - 2}"].font = SECTION_FONT
    header_row(sheet, top - 1, ["Monthly income band", "SKPG 2024, first degree (published)",
                                "Simulated sample", "Difference"])
    published = data_module.published_band_shares(skpg)
    simulated = data_module.sample_band_shares(frame["salary"])
    for offset, (label, real, drawn) in enumerate(zip(data_module.BAND_LABELS, published, simulated)):
        row = top + offset
        put(sheet, f"A{row}", label)
        put(sheet, f"B{row}", float(real), PERCENT)
        put(sheet, f"C{row}", float(drawn), PERCENT)
        put(sheet, f"D{row}", f"=C{row}-B{row}", PERCENT)

    widths(sheet, {"A": 36, **{get_column_letter(i): 20 for i in range(2, 15)}, "O": 90})


def build_notes(workbook, result):
    sheet = workbook.create_sheet(SHEET_NOTES)
    sheet["A1"] = "Notes"
    sheet["A1"].font = TITLE_FONT
    lines = [
        "What is real",
        "Salaries: drawn from the monthly income bands of first-degree graduates in KPT's Laporan Kajian "
        "Pengesanan Graduan 2024, Jadual 4o (n = 109,414), field by field.",
        "Policy figures: Budget 2027 (9 October 2026). Minimum wage RM1,700 to RM2,000 from June 2027; "
        "RM2,500 floor for graduates and semi-skilled workers, start date not announced.",
        "",
        "What is assumed",
        "CGPA, internship months and Klang Valley location are not published alongside income. They are "
        "simulated, and the strength of their link to salary is set by assumption.",
        "So the CGPA, internship, location and interaction coefficients show what the method recovers from "
        "those assumptions. They are not evidence about the Malaysian graduate labour market.",
        "The STEM coefficient is different: it comes from the real field-by-field salary bands, given this "
        "project's definition of STEM (engineering, ICT, natural sciences/mathematics/statistics).",
        "Salaries inside each band are spread evenly, the bottom band starts at RM500 and the top band is "
        "capped at RM15,000.",
        "",
        "Reading the model",
        "The dependent variable is ln(salary), so a coefficient b means roughly a 100 x (exp(b) - 1) per cent "
        "change in salary.",
        "CGPA is centred on the sample mean, so the STEM coefficient is the STEM gap at an average CGPA.",
        f"Predictions in RM multiply exp(prediction) by the smearing factor ({result['smearing']:.4f}); "
        "without it they would give the median, not the mean.",
        "Inference uses HC3 robust standard errors. "
        + ("The Breusch-Pagan or White test rejects constant error variance, so they are needed."
           if result["heteroscedastic"] else
           "Neither the Breusch-Pagan nor the White test rejects constant error variance at 5%; "
           "robust errors are valid either way."),
        "",
        "Limits",
        "Statutory contributions use flat rates. SOCSO and EIS really follow contribution tables with a wage "
        "ceiling, and income tax (PCB) is left out of take-home pay.",
        "The simulation assumes no job losses, no change in hours and full compliance among covered employers.",
        "The source data are 2024 salaries; scenario D on the Policy_Simulation tab shows the effect of pay "
        "growth before the floor starts.",
    ]
    for offset, line in enumerate(lines):
        cell = sheet.cell(row=3 + offset, column=1, value=line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if line in ("What is real", "What is assumed", "Reading the model", "Limits"):
            cell.font = SECTION_FONT
    widths(sheet, {"A": 130})


def build_workbook(result, scenarios, profile, parameters, skpg, path):
    workbook = Workbook()
    build_parameters(workbook, parameters, result["cgpa_mean"])
    build_data(workbook, result["frame"])
    build_ols(workbook, result)
    build_dashboard(workbook, result, parameters)
    build_policy(workbook, scenarios, profile, skpg, result["frame"])
    build_notes(workbook, result)
    workbook.active = workbook.sheetnames.index(SHEET_DASHBOARD)
    # openpyxl stores formulas without results; this makes Excel calculate on open.
    workbook.calculation.fullCalcOnLoad = True
    workbook.save(path)
    return path
